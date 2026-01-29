"""
Data export functionality for articles
"""

import json
import logging
from io import BytesIO
from typing import List, Dict, Optional
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .models import Article, get_db

logger = logging.getLogger(__name__)


async def export_articles(
    task_id: int,
    format: str = 'excel',
    filters: Optional[Dict] = None
) -> BytesIO:
    """
    Export scraped articles in various formats

    Args:
        task_id: Task ID to export articles from
        format: Output format (json, csv, excel)
        filters: Optional filters (start_date, end_date)

    Returns:
        BytesIO buffer with exported data
    """
    # Input validation
    if not isinstance(task_id, int) or task_id <= 0:
        raise ValueError(f"Invalid task_id: {task_id}. Must be a positive integer.")

    supported_formats = ['json', 'csv', 'excel']
    if format not in supported_formats:
        raise ValueError(f"Unsupported format: {format}. Must be one of {supported_formats}")

    # Validate filter dates if provided
    if filters:
        if 'start_date' in filters and not isinstance(filters['start_date'], datetime):
            raise ValueError("start_date must be a datetime object")
        if 'end_date' in filters and not isinstance(filters['end_date'], datetime):
            raise ValueError("end_date must be a datetime object")

    logger.info(f"Exporting articles for task {task_id} as {format}")

    # Query articles with proper resource management
    db = next(get_db())
    try:
        query = db.query(Article).filter(Article.task_id == task_id)

        # Apply filters
        if filters:
            if 'start_date' in filters:
                query = query.filter(Article.publish_time >= filters['start_date'])
            if 'end_date' in filters:
                query = query.filter(Article.publish_time <= filters['end_date'])

        articles = query.order_by(Article.publish_time.desc()).all()
        logger.info(f"Found {len(articles)} articles to export")
    except Exception as e:
        logger.error(f"Database query failed for task {task_id}: {e}")
        raise
    finally:
        db.close()

    # Export with error handling
    try:
        if format == 'json':
            return export_as_json(articles)
        elif format == 'csv':
            return export_as_csv(articles)
        elif format == 'excel':
            return export_as_excel(articles)
    except Exception as e:
        logger.error(f"Export failed for format {format}: {e}")
        raise


def export_as_excel(articles: List[Article]) -> BytesIO:
    """Export to Excel with formatting"""
    logger.info("Exporting as Excel")

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Articles"

        # Headers
        headers = [
            'ID', 'Title', 'Author', 'Publish Time',
            'URL', 'Word Count', 'Crawled At'
        ]

        # Style headers
        header_fill = PatternFill(
            start_color='4472C4',
            end_color='4472C4',
            fill_type='solid'
        )
        header_font = Font(bold=True, color='FFFFFF')

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # Data rows
        for row, article in enumerate(articles, 2):
            word_count = len(article.content.split()) if article.content else 0

            ws.cell(row, 1, article.id)
            ws.cell(row, 2, article.title)
            ws.cell(row, 3, article.author)
            ws.cell(row, 4, article.publish_time.strftime('%Y-%m-%d %H:%M') if article.publish_time else '')
            ws.cell(row, 5, article.url)
            ws.cell(row, 6, word_count)
            ws.cell(row, 7, article.crawled_at.strftime('%Y-%m-%d %H:%M') if article.crawled_at else '')

        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

        # Add summary sheet
        summary = wb.create_sheet("Summary")
        summary['A1'] = 'Total Articles'
        summary['B1'] = len(articles)
        summary['A2'] = 'Export Time'
        summary['B2'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

        # Save to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return buffer
    except Exception as e:
        logger.error(f"Excel export failed: {e}")
        raise


def export_as_csv(articles: List[Article]) -> BytesIO:
    """Export to CSV format"""
    logger.info("Exporting as CSV")

    try:
        df = pd.DataFrame([
            {
                'ID': a.id,
                'Title': a.title,
                'Author': a.author,
                'Publish Time': a.publish_time.strftime('%Y-%m-%d %H:%M') if a.publish_time else '',
                'URL': a.url,
                'Content Preview': (a.content[:200] + '...' if len(a.content) > 200 else a.content) if a.content else '',
                'Crawled At': a.crawled_at.strftime('%Y-%m-%d %H:%M') if a.crawled_at else ''
            }
            for a in articles
        ])

        buffer = BytesIO()
        df.to_csv(buffer, index=False, encoding='utf-8-sig')
        buffer.seek(0)

        return buffer
    except Exception as e:
        logger.error(f"CSV export failed: {e}")
        raise


def export_as_json(articles: List[Article]) -> BytesIO:
    """Export to JSON format"""
    logger.info("Exporting as JSON")

    try:
        data = {
            'export_time': datetime.now().isoformat(),
            'total_articles': len(articles),
            'articles': [a.to_dict() for a in articles]
        }

        buffer = BytesIO()
        buffer.write(json.dumps(data, ensure_ascii=False, indent=2).encode('utf-8'))
        buffer.seek(0)

        return buffer
    except Exception as e:
        logger.error(f"JSON export failed: {e}")
        raise
